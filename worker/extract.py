"""חילוץ טקסט מ-PDF ומ-DOCX. ראו docs/10 §10.2."""

from pathlib import Path

from worker.chunking import clean_text


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return clean_text(_pdf(path))
    if suffix == ".docx":
        return clean_text(_docx(path))
    raise ValueError("נדרש מסמך PDF או DOCX")


def _pdf(path: Path) -> str:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(str(path))
    except PdfReadError as exc:
        raise ValueError("חילוץ הטקסט נכשל") from exc
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _docx(path: Path) -> str:
    from docx import Document
    from docx.opc.exceptions import PackageNotFoundError

    try:
        document = Document(str(path))
    except PackageNotFoundError as exc:
        raise ValueError("חילוץ הטקסט נכשל") from exc
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)
