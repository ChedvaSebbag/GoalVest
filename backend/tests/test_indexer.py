"""Indexer עם Chroma ו-Ollama מדומים, כולל הרצה חוזרת בלי כפילויות."""

import io
from decimal import Decimal

import pytest
import sqlalchemy as sa
from app.bootstrap import build_dispatcher
from app.infrastructure.event_store import EventStoreRepository
from app.infrastructure.tables import consumer_checkpoints, event_store, rm_product_documents
from app.products.storage import write_stored
from app.queries.products import list_products
from app.reference.institutions import INSTITUTIONS
from docx import Document
from goalvest_core.products import Product

from worker.index_knowledge import index_knowledge
from worker.indexer import run_once
from worker.rag import search_knowledge_base

INSTITUTION = INSTITUTIONS[0]


class FakeCollection:
    def __init__(self):
        self.upserts: list[list[str]] = []
        self.rows: dict[str, str] = {}
        self.metadatas: dict[str, dict] = {}

    def upsert(self, ids, embeddings, documents, metadatas):
        self.upserts.append(list(ids))
        for index, item_id in enumerate(ids):
            self.rows[item_id] = documents[index]
            self.metadatas[item_id] = metadatas[index]

    def query(self, **kwargs):
        self.last_query = kwargs
        return {
            "documents": [["קטע על נזילות"]],
            "metadatas": [[{"title": "תנאי פוליסה", "document_id": "doc-1"}]],
            "distances": [[0.25]],
        }


def _embed(texts: list[str]) -> list[list[float]]:
    return [[float(len(text)), 0.5] for text in texts]


def _docx(text: str) -> bytes:
    buffer = io.BytesIO()
    document = Document()
    if text:
        document.add_paragraph(text)
    document.save(buffer)
    return buffer.getvalue()


def _fields():
    return {
        "product_type": "HISHTALMUT",
        "name": "אופק השתלמות",
        "track": "כללי",
        "symbol": None,
        "fee_deposit": Decimal("0.00"),
        "fee_assets": Decimal("0.55"),
        "return_1y": Decimal("8.00"),
        "return_3y": Decimal("7.00"),
        "return_5y": Decimal("6.00"),
        "std_dev": Decimal("5.00"),
        "risk_level": 3,
        "liquidity": "PARTIAL",
        "liquidity_terms": "נזילות חלקית אחרי תקופת הנעילה.",
        "tax_notes": "כללי המס לפי הדין העדכני.",
        "min_deposit": Decimal("0.00"),
        "description": "מסלול להמחשה.",
    }


def _upload(session, tmp_path, text: str):
    repository = EventStoreRepository(session, build_dispatcher())
    product = Product()
    product.create(institution_id=INSTITUTION.id, known_institutions={str(INSTITUTION.id)}, **_fields())
    repository.append(product.aggregate_id, 0, product.uncommitted_events)
    session.commit()
    current = Product.from_events(repository.load(product.aggregate_id, "Product"))
    content = _docx(text)
    current.upload_document(content=content, original_name="כללים.docx", doc_type="TAX_RULES")
    stored = current.uncommitted_events[-1].payload
    write_stored(tmp_path, stored["stored_name"], content)
    repository.append(current.aggregate_id, current.version - 1, current.uncommitted_events)
    session.commit()
    return stored["document_id"]


def _statuses(session) -> list[str]:
    return list(
        session.execute(sa.select(event_store.c.event_type).order_by(event_store.c.global_position)).scalars()
    )


def test_retry_after_stopped_commit_upserts_the_same_ids(session, tmp_path):
    document_id = _upload(session, tmp_path, "מיסוי " * 200)
    collection = FakeCollection()
    state = {"failed": False}

    def commit(db):
        if not state["failed"]:
            state["failed"] = True
            db.rollback()
            raise RuntimeError("stopped")
        db.commit()

    with pytest.raises(RuntimeError, match="stopped"):
        run_once(
            session,
            storage_dir=tmp_path,
            embedder=_embed,
            collection=collection,
            model="bge-m3",
            commit=commit,
        )
    assert "ProductDocumentIndexed" not in _statuses(session)
    assert session.execute(sa.select(consumer_checkpoints)).first() is None
    assert collection.upserts

    run_once(session, storage_dir=tmp_path, embedder=_embed, collection=collection, model="bge-m3", commit=commit)
    assert collection.upserts[0] == collection.upserts[1]
    assert collection.upserts[0] == [f"{document_id}:{index}" for index in range(len(collection.upserts[0]))]
    assert len(collection.rows) == len(collection.upserts[0])
    assert _statuses(session).count("ProductDocumentIndexed") == 1
    document = session.execute(sa.select(rm_product_documents)).mappings().one()
    assert document["index_status"] == "INDEXED"
    assert document["chunks_count"] == len(collection.rows)
    position = session.execute(
        sa.select(event_store.c.global_position).where(event_store.c.event_type == "ProductDocumentUploaded")
    ).scalar_one()
    checkpoint = session.execute(sa.select(consumer_checkpoints.c.last_position)).scalar_one()
    assert checkpoint == position
    page = list_products(
        session,
        include_inactive=True,
        viewer_role="AGENT",
        viewer_institution=INSTITUTION.id,
    )
    assert page["items"][0]["documents"][0]["indexStatus"] == "INDEXED"


def test_empty_document_is_marked_failed_without_upsert(session, tmp_path):
    _upload(session, tmp_path, "")
    collection = FakeCollection()
    assert run_once(session, storage_dir=tmp_path, embedder=_embed, collection=collection, model="bge-m3") == 1
    assert collection.upserts == []
    assert _statuses(session).count("ProductDocumentIndexFailed") == 1
    document = session.execute(sa.select(rm_product_documents)).mappings().one()
    assert document["index_status"] == "INDEX_FAILED"
    assert session.execute(sa.select(consumer_checkpoints.c.last_position)).scalar_one() > 0


def test_knowledge_upsert_uses_document_and_chunk_ids(tmp_path):
    folder = tmp_path / "knowledge"
    folder.mkdir()
    note = """---
title: קרן חירום
topic: liquidity
source_url: seed://knowledge/emergency-fund
valid_from: 2026-01-01
---

קרן חירום נזילה.
"""
    (folder / "emergency-fund.md").write_text(note, encoding="utf-8")
    collection = FakeCollection()
    assert index_knowledge(folder, collection, _embed) == 1
    assert collection.upserts == [["emergency-fund:0"]]
    assert collection.metadatas["emergency-fund:0"]["topic"] == "liquidity"


def test_search_returns_title_excerpt_and_score():
    store = FakeCollection()
    found = search_knowledge_base(
        "נזילות",
        "policy_documents",
        k=6,
        product_ids=["product-1"],
        store=store,
        embedder=_embed,
    )
    assert store.last_query["where"] == {"product_id": {"$in": ["product-1"]}}
    assert found == [{"title": "תנאי פוליסה", "document_id": "doc-1", "excerpt": "קטע על נזילות", "score": 0.75}]
